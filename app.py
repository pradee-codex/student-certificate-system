from flask import (
    Flask,
    render_template,
    request,
    redirect,
    session,
    flash,
    send_from_directory,
    send_file,
    url_for
)

from flask_mysqldb import MySQL
from config import Config
from flask import send_from_directory
from werkzeug.utils import secure_filename
from drive_service import upload_to_drive

import os
import zipfile
import tempfile
import pandas as pd

# ==========================
# NEW IMPORTS
# ==========================

from openpyxl import Workbook
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from io import BytesIO

app = Flask(__name__)

app.config.from_object(Config)
app.secret_key = Config.SECRET_KEY

# ==========================
# Startup diagnostics:
# logs whether the Drive service account key is visible
# ==========================

import logging

logging.basicConfig(level=logging.INFO)

_drive_key_paths = [
    os.getenv("GOOGLE_SERVICE_ACCOUNT_FILE"),
    "/etc/secrets/key.json",
    os.path.join(os.path.dirname(os.path.abspath(__file__)), "key.json")
]

if not any(p and os.path.exists(p) for p in _drive_key_paths):
    logging.warning(
        "Drive service account key NOT found. Checked: %s",
        [p for p in _drive_key_paths if p]
    )
else:
    for _p in _drive_key_paths:
        if _p and os.path.exists(_p):
            logging.info("Drive service account key FOUND at: %s", _p)

# ==========================
# Upload Folder
# ==========================

UPLOAD_FOLDER = "static/uploads/certificates"

app.config["UPLOAD_FOLDER"] = UPLOAD_FOLDER

if not os.path.exists(UPLOAD_FOLDER):
    os.makedirs(UPLOAD_FOLDER)

UPLOAD_FOLDER = "uploads/excel"
os.makedirs(UPLOAD_FOLDER, exist_ok=True)

PROFILE_FOLDER = "static/uploads/profile"
app.config["PROFILE_FOLDER"] = PROFILE_FOLDER
os.makedirs(PROFILE_FOLDER, exist_ok=True)

# ==========================
# MySQL
# ==========================


app.config["MYSQL_HOST"] = os.getenv("MYSQL_HOST", "").strip()
app.config["MYSQL_USER"] = os.getenv("MYSQL_USER", "").strip()
app.config["MYSQL_PASSWORD"] = os.getenv("MYSQL_PASSWORD", "").strip()
app.config["MYSQL_DB"] = os.getenv("MYSQL_DB", "").strip()
app.config["MYSQL_PORT"] = int(os.getenv("MYSQL_PORT", "3306").strip())

app.config["MYSQL_SSL"] = {
    "ssl": {}
}

mysql = MySQL(app)


@app.route("/test-db")
def test_db():

    try:
        cursor = mysql.connection.cursor()
        cursor.execute("SELECT 1")
        result = cursor.fetchone()
        cursor.close()

        return f"Database connected successfully: {result}"

    except Exception as e:
             return f"Database ERROR: {str(e)}", 500

# ==========================
# Home
# ==========================

@app.route("/")
def home():
    return render_template("auth/login.html")

# ==========================
# Login
# ==========================
@app.route("/login", methods=["GET","POST"])
def login():

    username = request.form["username"]
    password = request.form["password"]

    cursor = mysql.connection.cursor()

    cursor.execute("""
    SELECT
        id,
        username,
        role,
        profile_photo
    FROM users
    WHERE username=%s
    AND password=%s
""", (username, password))

    user = cursor.fetchone()

    cursor.close()

    if user:

        session["user_id"] = user[0]
        session["username"] = user[1]
        session["role"] = user[2]
        session["profile_photo"] = user[3]
        if user[2] == "admin":
            return redirect("/admin")

        elif user[2] == "student":
            return redirect("/student")

        elif user[2] == "tutor":
            return redirect("/tutor")
        
        elif user[2] == "hod":
            return redirect("/hod")

    flash("Invalid Username or Password")
    return redirect("/")

    # ==========================
# Admin Dashboard
# ==========================

@app.route("/admin")
def admin():

    if session.get("role") != "admin":
        return redirect("/")

    cursor = mysql.connection.cursor()

    # ==============================
    # Dashboard Count
    # ==============================

    cursor.execute("SELECT COUNT(*) FROM students")
    total_students = cursor.fetchone()[0]

    cursor.execute("SELECT COUNT(*) FROM certificates")
    total_certificates = cursor.fetchone()[0]

    # ==============================
    # Student List
    # ==============================

    cursor.execute("""
        SELECT
            student_id,
            register_no,
            student_name,
            department,
            year,
            email
        FROM students
        ORDER BY student_name
    """)

    students = cursor.fetchall()

    # ==============================
    # Certificate List
    # ==============================

    cursor.execute("""
        SELECT
            students.student_name,
            students.department,
            students.year,
            certificates.certificate_title,
            certificate_categories.category_name,
            certificates.achievement,
            certificates.certificate_file,
            certificates.upload_date

        FROM certificates

        INNER JOIN students
            ON certificates.student_id = students.student_id

        INNER JOIN certificate_categories
            ON certificates.category_id =
               certificate_categories.category_id

        ORDER BY certificates.upload_date DESC
    """)

    certificates = cursor.fetchall()

    cursor.close()

    return render_template(
        "admin/dashboard.html",
        total_students=total_students,
        total_certificates=total_certificates,
        students=students,
        certificates=certificates
    )

#-------------------------------------------
# ==========================================
# ADMIN PROFILE
# ==========================================

@app.route("/admin_profile")
def admin_profile():

    if session.get("role") != "admin":
        return redirect("/")

    cursor = mysql.connection.cursor()

    cursor.execute("""
        SELECT
            id,
            username,
            full_name,
            email,
            phone,
            profile_photo
        FROM users
        WHERE id=%s
    """, (session["user_id"],))

    admin = cursor.fetchone()

    cursor.close()

    return render_template("admin/admin_profile.html", admin=admin)

@app.route("/update_admin_profile", methods=["POST"])
def update_admin_profile():

    if session.get("role") != "admin":
        return redirect("/")

    full_name = request.form["full_name"]
    email = request.form["email"]
    phone = request.form["phone"]

    cursor = mysql.connection.cursor()

    cursor.execute("""
        UPDATE users
        SET
            full_name=%s,
            email=%s,
            phone=%s
        WHERE id=%s
    """, (
        full_name,
        email,
        phone,
        session["user_id"]
    ))

    mysql.connection.commit()

    cursor.close()

    flash("Profile Updated Successfully")

    return redirect("/admin_profile")


# ==========================
# Reports
# ==========================

@app.route("/reports")
def reports():

    if session.get("role") != "admin":
        return redirect("/")

    cursor = mysql.connection.cursor()

    # Total Students
    cursor.execute("SELECT COUNT(*) FROM students")
    total_students = cursor.fetchone()[0]

    # Total Certificates
    cursor.execute("SELECT COUNT(*) FROM certificates")
    total_certificates = cursor.fetchone()[0]

    # Department Report
    cursor.execute("""
        SELECT department, COUNT(*)
        FROM students
        GROUP BY department
        ORDER BY department
    """)
    department_report = cursor.fetchall()

    # Year Report
    cursor.execute("""
        SELECT year, COUNT(*)
        FROM students
        GROUP BY year
        ORDER BY year
    """)
    year_report = cursor.fetchall()

    # Category Report
    cursor.execute("""
        SELECT
            certificate_categories.category_name,
            COUNT(certificates.certificate_id)

        FROM certificate_categories

        LEFT JOIN certificates
        ON certificate_categories.category_id = certificates.category_id

        GROUP BY certificate_categories.category_name

        ORDER BY certificate_categories.category_name
    """)
    category_report = cursor.fetchall()

    # Certificate Report
    cursor.execute("""
        SELECT
            students.register_no,
            students.student_name,
            students.department,
            students.year,
            certificates.certificate_title,
            certificate_categories.category_name,
            certificates.achievement,
            certificates.certificate_file,
            certificates.upload_date

        FROM certificates

        INNER JOIN students
            ON certificates.student_id = students.student_id

        INNER JOIN certificate_categories
            ON certificates.category_id = certificate_categories.category_id

        ORDER BY certificates.upload_date DESC
    """)

    certificates = cursor.fetchall()

    cursor.close()

    return render_template(
        "admin/report.html",
        total_students=total_students,
        total_certificates=total_certificates,
        department_report=department_report,
        year_report=year_report,
        category_report=category_report,
        certificates=certificates
    )

# ==========================
# Manage HOD
# ==========================

@app.route("/manage_hod")
def manage_hod():

    if session.get("role") != "admin":
        return redirect("/")

    cursor = mysql.connection.cursor()

    cursor.execute("""
        SELECT
            hods.hod_id,
            hods.hod_name,
            hods.department,
            hods.email,
            hods.phone,
            users.username

        FROM hods

        INNER JOIN users
        ON hods.user_id = users.id

        ORDER BY hod_name
    """)

    hods = cursor.fetchall()

    cursor.close()

    return render_template(
        "admin/manage_hod.html",
        hods=hods
    )

# ==========================
# Add HOD
# ==========================

@app.route("/add_hod", methods=["POST"])
def add_hod():

    if session.get("role") != "admin":
        return redirect("/")

    hod_name = request.form["hod_name"]
    department = request.form["department"]
    email = request.form["email"]
    phone = request.form["phone"]
    username = request.form["username"]
    password = request.form["password"]

    cursor = mysql.connection.cursor()

    # Username already exists?
    cursor.execute(
        "SELECT id FROM users WHERE username=%s",
        (username,)
    )

    existing = cursor.fetchone()

    if existing:
        cursor.close()
        flash("Username already exists")
        return redirect("/manage_hod")

    # Insert into users table
    cursor.execute("""
        INSERT INTO users
        (
            username,
            password,
            role
        )
        VALUES
        (
            %s,
            %s,
            'hod'
        )
    """, (
        username,
        password
    ))

    mysql.connection.commit()

    user_id = cursor.lastrowid

    # Insert into hods table
    cursor.execute("""
        INSERT INTO hods
        (
            user_id,
            hod_name,
            department,
            email,
            phone
        )
        VALUES
        (
            %s,
            %s,
            %s,
            %s,
            %s
        )
    """, (
        user_id,
        hod_name,
        department,
        email,
        phone
    ))

    mysql.connection.commit()

    cursor.close()

    flash("HOD Created Successfully")

    return redirect("/manage_hod")
#-------------------------------

@app.route("/admin_export_zip")
def admin_export_zip():

    if session.get("role") != "admin":
        return redirect("/")

    search_text = request.args.get("search", "").strip()

    cursor = mysql.connection.cursor()

    # ==========================================
    # GET CERTIFICATES
    # ==========================================

    query = """
        SELECT
            students.student_name,
            students.department,
            students.year,
            certificates.certificate_title,
            certificate_categories.category_name,
            certificates.achievement,
            certificates.upload_date,
            certificates.certificate_file

        FROM certificates

        INNER JOIN students
            ON certificates.student_id = students.student_id

        INNER JOIN certificate_categories
            ON certificates.category_id = certificate_categories.category_id

        WHERE 1=1
    """

    params = []

    # ==========================================
    # SEARCH FILTER
    # ==========================================

    if search_text:

        query += """
            AND (
                students.student_name LIKE %s
                OR students.department LIKE %s
                OR certificates.certificate_title LIKE %s
                OR students.register_no LIKE %s
            )
        """

        search_pattern = "%" + search_text + "%"

        params.extend([
            search_pattern,
            search_pattern,
            search_pattern,
            search_pattern
        ])

    # ==========================================
    # ORDER
    # ==========================================

    query += """
        ORDER BY certificates.upload_date DESC
    """

    cursor.execute(query, tuple(params))

    certificates = cursor.fetchall()

    cursor.close()

    # ==========================================
    # NO DATABASE RESULT
    # ==========================================

    if not certificates:

        flash("No certificates found for the search.")

        return redirect("/admin")

    # ==========================================
    # CERTIFICATE FOLDER
    # ==========================================

    cert_folder = os.path.join(
        app.root_path,
        "static",
        "uploads",
        "certificates"
    )

    # ==========================================
    # CREATE TEMP ZIP
    # ==========================================

    temp_zip = tempfile.NamedTemporaryFile(
        delete=False,
        suffix=".zip"
    )

    zip_path = temp_zip.name

    temp_zip.close()

    added_files = 0

    # ==========================================
    # CREATE ZIP
    # ==========================================

    with zipfile.ZipFile(
        zip_path,
        "w",
        zipfile.ZIP_DEFLATED
    ) as zipf:

        for certificate in certificates:

            # certificate_file is index 7
            filename = certificate[7]

            if not filename:
                continue

            file_path = os.path.join(
                cert_folder,
                filename
            )

            if os.path.isfile(file_path):

                zipf.write(
                    file_path,
                    arcname=filename
                )

                added_files += 1

    # ==========================================
    # NO PHYSICAL FILE
    # ==========================================

    if added_files == 0:

        os.remove(zip_path)

        flash("Certificate files not found.")

        return redirect("/admin")

    # ==========================================
    # ZIP NAME
    # ==========================================

    if search_text:

        download_name = "Admin_Search_Certificates.zip"

    else:

        download_name = "Admin_All_Certificates.zip"

    # ==========================================
    # SEND ZIP
    # ==========================================

    return send_file(
        zip_path,
        as_attachment=True,
        download_name=download_name,
        mimetype="application/zip"
    )

# ==========================
# Tutor Dashboard
# ==========================


#------------------------------------
@app.route("/tutor")
def tutor():

    if session.get("role") != "tutor":
        return redirect("/")

    cursor = mysql.connection.cursor()

    # =========================================================
    # Tutor Details
    # =========================================================

    cursor.execute("""
        SELECT
            tutor_name,
            department,
            class_year,
            section,
            profile_photo
        FROM tutors
        WHERE user_id=%s
    """, (session["user_id"],))

    tutor = cursor.fetchone()

    if not tutor:
        cursor.close()
        flash("Tutor not assigned.")
        return redirect("/logout")

    tutor_name = tutor[0]
    department = tutor[1]
    class_year = tutor[2]
    section = tutor[3]

    print("====================================")
    print("LOGIN USER ID:", session["user_id"])
    print("TUTOR DATA:", tutor)
    print("DEPARTMENT:", department)
    print("CLASS YEAR:", class_year)
    print("SECTION:", section)
    print("====================================")


    # =========================================================
    # COMMON STUDENT FILTER
    # =========================================================

    if class_year is None or section is None:

        student_where = """
            students.department=%s
        """

        student_params = (
            department,
        )

    else:

        student_where = """
            students.department=%s
            AND students.year=%s
            AND students.section=%s
        """

        student_params = (
            department,
            class_year,
            section
        )


    # =========================================================
    # DASHBOARD - TOTAL STUDENTS
    # =========================================================

    cursor.execute(f"""
        SELECT COUNT(*)
        FROM students
        WHERE {student_where}
    """, student_params)

    total_students = cursor.fetchone()[0]


    # =========================================================
    # DASHBOARD - TOTAL CERTIFICATES
    # =========================================================

    cursor.execute(f"""
        SELECT COUNT(*)
        FROM certificates

        INNER JOIN students
            ON certificates.student_id = students.student_id

        WHERE {student_where}
    """, student_params)

    total_certificates = cursor.fetchone()[0]


    # =========================================================
    # STUDENT LIST
    # =========================================================

    cursor.execute(f"""
        SELECT
            students.student_name,
            students.register_no,
            students.department,
            students.year,
            students.email

        FROM students

        WHERE {student_where}

        ORDER BY students.student_name
    """, student_params)

    students = cursor.fetchall()


    # =========================================================
    # CERTIFICATE LIST
    # =========================================================

    cursor.execute(f"""
        SELECT
            students.student_name,
            students.department,
            students.year,
            certificates.certificate_title,
            certificate_categories.category_name,
            certificates.achievement,
            students.profile_photo,
            certificates.upload_date,
            certificates.certificate_file

        FROM certificates

        INNER JOIN students
            ON certificates.student_id = students.student_id

        INNER JOIN certificate_categories
            ON certificates.category_id =
               certificate_categories.category_id

        WHERE {student_where}

        ORDER BY certificates.upload_date DESC
    """, student_params)

    certificates = cursor.fetchall()


    # =========================================================
    # STUDENT SEARCH + CHART DATA
    # =========================================================
    #
    # Used for:
    # Search Student
    # Total Certificates
    # Winner
    # Participated
    # Certificate Category Chart
    # Achievement Type Chart
    #
    # =========================================================

    # Student name -> Register number
    student_lookup = {
        str(student[0]).strip().lower(): student[1]
        for student in students
    }

    student_chart_data = []

    for row in certificates:

        student_name = row[0] if row[0] else ""

        student_chart_data.append({

            "name": student_name,

            "register_no": student_lookup.get(
                str(student_name).strip().lower(),
                ""
            ),

            "department": row[1] if row[1] else "",

            "year": row[2] if row[2] else "",

            "certificate": row[3] if row[3] else "",

            "category": row[4] if row[4] else "Others",

            "achievement": row[5] if row[5] else "Others"
        })


    # =========================================================
    # STUDENT-WISE CERTIFICATE COUNT
    # =========================================================

    cursor.execute(f"""
        SELECT
            students.student_name,
            students.register_no,
            students.department,
            COUNT(certificates.student_id) AS total_certificates

        FROM students

        LEFT JOIN certificates
            ON students.student_id = certificates.student_id

        WHERE {student_where}

        GROUP BY
            students.student_id,
            students.student_name,
            students.register_no,
            students.department

        ORDER BY total_certificates DESC
    """, student_params)

    student_certificate_counts = cursor.fetchall()


    # =========================================================
    # TOP STUDENTS
    # =========================================================

    top_students = student_certificate_counts[:10]


    # =========================================================
    # CERTIFICATE CATEGORY DISTRIBUTION
    # =========================================================

    cursor.execute(f"""
        SELECT
            certificate_categories.category_name,
            COUNT(certificates.certificate_id)

        FROM certificates

        INNER JOIN students
            ON certificates.student_id = students.student_id

        INNER JOIN certificate_categories
            ON certificates.category_id =
               certificate_categories.category_id

        WHERE {student_where}

        GROUP BY
            certificate_categories.category_id,
            certificate_categories.category_name

        ORDER BY COUNT(certificates.certificate_id) DESC
    """, student_params)

    category_report = cursor.fetchall()


    # =========================================================
    # ACHIEVEMENT TYPE DISTRIBUTION
    # =========================================================

    cursor.execute(f"""
        SELECT
            certificates.achievement,
            COUNT(certificates.certificate_id)

        FROM certificates

        INNER JOIN students
            ON certificates.student_id = students.student_id

        WHERE {student_where}

        GROUP BY certificates.achievement

        ORDER BY COUNT(certificates.certificate_id) DESC
    """, student_params)

    achievement_report = cursor.fetchall()


    # =========================================================
    # STUDENT + CATEGORY REPORT
    # =========================================================

    cursor.execute(f"""
        SELECT
            students.student_name,
            certificate_categories.category_name,
            COUNT(certificates.certificate_id)

        FROM certificates

        INNER JOIN students
            ON certificates.student_id = students.student_id

        INNER JOIN certificate_categories
            ON certificates.category_id =
               certificate_categories.category_id

        WHERE {student_where}

        GROUP BY
            students.student_id,
            students.student_name,
            certificate_categories.category_id,
            certificate_categories.category_name

        ORDER BY
            students.student_name
    """, student_params)

    student_category_report = cursor.fetchall()


    # =========================================================
    # STUDENT + ACHIEVEMENT REPORT
    # =========================================================

    cursor.execute(f"""
        SELECT
            students.student_name,
            certificates.achievement,
            COUNT(certificates.certificate_id)

        FROM certificates

        INNER JOIN students
            ON certificates.student_id = students.student_id

        WHERE {student_where}

        GROUP BY
            students.student_id,
            students.student_name,
            certificates.achievement

        ORDER BY
            students.student_name
    """, student_params)

    student_achievement_report = cursor.fetchall()


    # =========================================================
    # DEBUG
    # =========================================================

    print("====================================")
    print("TOTAL STUDENTS:", total_students)
    print("TOTAL CERTIFICATES:", total_certificates)
    print(
        "STUDENT CERTIFICATE COUNTS:",
        student_certificate_counts
    )
    print("CATEGORY REPORT:", category_report)
    print("ACHIEVEMENT REPORT:", achievement_report)
    print("STUDENT CHART DATA:", student_chart_data)
    print("====================================")


    # =========================================================
    # CLOSE CURSOR
    # =========================================================

    cursor.close()


    # =========================================================
    # RENDER DASHBOARD
    # =========================================================

    return render_template(
        "tutor/dashboard.html",

        # -----------------------------------------------------
        # Tutor Details
        # -----------------------------------------------------

        tutor=tutor,
        tutor_name=tutor_name,
        department=department,
        class_year=class_year,
        section=section,


        # -----------------------------------------------------
        # Dashboard
        # -----------------------------------------------------

        total_students=total_students,
        total_certificates=total_certificates,


        # -----------------------------------------------------
        # Student / Certificate Lists
        # -----------------------------------------------------

        students=students,
        certificates=certificates,


        # -----------------------------------------------------
        # Existing Graph / Report Data
        # -----------------------------------------------------

        top_students=top_students,

        student_certificate_counts=
            student_certificate_counts,

        category_report=
            category_report,

        achievement_report=
            achievement_report,

        student_category_report=
            student_category_report,

        student_achievement_report=
            student_achievement_report,


        # -----------------------------------------------------
        # NEW
        # Student Search + Category + Achievement Charts
        # -----------------------------------------------------

        student_chart_data=
            student_chart_data
    )
#=====================================================
    

@app.route("/tutor_profile")
def tutor_profile():

    if session.get("role") != "tutor":
        return redirect("/")

    cursor = mysql.connection.cursor()

    cursor.execute("""
        SELECT
            tutors.tutor_id,
            tutors.tutor_name,
            tutors.department,
            tutors.class_year,
            tutors.section,
            tutors.email,
            tutors.phone,
            users.username,
            tutors.profile_photo

        FROM tutors

        INNER JOIN users
            ON tutors.user_id = users.id

        WHERE users.id=%s
    """, (session["user_id"],))

    tutor = cursor.fetchone()

    cursor.close()

    return render_template(
        "tutor/profile.html",
        tutor=tutor
    )

@app.route("/update_tutor_profile", methods=["POST"])
def update_tutor_profile():

    if session.get("role") != "tutor":
        return redirect("/")

    name = request.form["name"]
    email = request.form["email"]
    phone = request.form["phone"]
    class_year = request.form["class_year"]
    section = request.form["section"]

    cursor = mysql.connection.cursor()

    cursor.execute("""
        UPDATE tutors
        SET
            tutor_name=%s,
            class_year=%s,
            section=%s,
            email=%s,
            phone=%s
        WHERE user_id=%s
    """, (
        name,
        class_year,
        section,
        email,
        phone,
        session["user_id"]
    ))

    mysql.connection.commit()

    cursor.close()

    flash("Profile Updated Successfully")

    return redirect("/tutor_profile")

#---------------------------
#       report
#----------------------------
@app.route("/tutor_reports")
def tutor_reports():

    if session.get("role") != "tutor":
        return redirect("/")

    cursor = mysql.connection.cursor()

    # =========================================================
    # Tutor Details
    # =========================================================

    cursor.execute("""
        SELECT
            tutor_name,
            department,
            class_year,
            section,
            profile_photo
        FROM tutors
        WHERE user_id=%s
    """, (session["user_id"],))

    tutor = cursor.fetchone()

    if not tutor:
        cursor.close()
        return redirect("/tutor")

    tutor_name = tutor[0]
    department = tutor[1]
    class_year = tutor[2]
    section = tutor[3]


    # =========================================================
    # Total Students
    # =========================================================

    if class_year is None or section is None:

        # No year / section
        # Count all students from tutor department

        cursor.execute("""
            SELECT COUNT(*)
            FROM students
            WHERE department=%s
        """, (department,))

    else:

        # Year + section assigned

        cursor.execute("""
            SELECT COUNT(*)
            FROM students
            WHERE department=%s
            AND year=%s
            AND section=%s
        """, (department, class_year, section))

    total_students = cursor.fetchone()[0]


    # =========================================================
    # Total Certificates
    # =========================================================

    if class_year is None or section is None:

        # All certificates from department

        cursor.execute("""
            SELECT COUNT(*)
            FROM certificates

            INNER JOIN students
                ON certificates.student_id =
                   students.student_id

            WHERE students.department=%s
        """, (department,))

    else:

        # Certificates from assigned year + section

        cursor.execute("""
            SELECT COUNT(*)
            FROM certificates

            INNER JOIN students
                ON certificates.student_id =
                   students.student_id

            WHERE students.department=%s
            AND students.year=%s
            AND students.section=%s
        """, (department, class_year, section))

    total_certificates = cursor.fetchone()[0]


    # =========================================================
    # Department Report
    # =========================================================

    if class_year is None or section is None:

        cursor.execute("""
            SELECT
                department,
                COUNT(*)
            FROM students

            WHERE department=%s

            GROUP BY department
        """, (department,))

    else:

        cursor.execute("""
            SELECT
                department,
                COUNT(*)
            FROM students

            WHERE department=%s
            AND year=%s
            AND section=%s

            GROUP BY department
        """, (department, class_year, section))

    department_report = cursor.fetchall()


    # =========================================================
    # Year Report
    # =========================================================

    if class_year is None or section is None:

        # Show year-wise students from department

        cursor.execute("""
            SELECT
                year,
                COUNT(*)
            FROM students

            WHERE department=%s

            GROUP BY year
            ORDER BY year
        """, (department,))

    else:

        cursor.execute("""
            SELECT
                year,
                COUNT(*)
            FROM students

            WHERE department=%s
            AND year=%s
            AND section=%s

            GROUP BY year
            ORDER BY year
        """, (department, class_year, section))

    year_report = cursor.fetchall()


    # =========================================================
    # Certificate Report
    # =========================================================

    if class_year is None or section is None:

        # ALL certificates from tutor department

        cursor.execute("""
            SELECT
                students.student_name,
                students.department,
                students.year,
                certificates.certificate_title,
                certificate_categories.category_name,
                certificates.achievement,
                certificates.certificate_file,
                certificates.upload_date

            FROM certificates

            INNER JOIN students
                ON certificates.student_id =
                   students.student_id

            INNER JOIN certificate_categories
                ON certificates.category_id =
                   certificate_categories.category_id

            WHERE students.department=%s

            ORDER BY certificates.upload_date DESC
        """, (department,))

    else:

        # Certificates from tutor's year + section

        cursor.execute("""
            SELECT
                students.student_name,
                students.department,
                students.year,
                certificates.certificate_title,
                certificate_categories.category_name,
                certificates.achievement,
                certificates.certificate_file,
                certificates.upload_date

            FROM certificates

            INNER JOIN students
                ON certificates.student_id =
                   students.student_id

            INNER JOIN certificate_categories
                ON certificates.category_id =
                   certificate_categories.category_id

            WHERE students.department=%s
            AND students.year=%s
            AND students.section=%s

            ORDER BY certificates.upload_date DESC
        """, (department, class_year, section))

    certificates = cursor.fetchall()


    # =========================================================
    # Debug
    # =========================================================

    print("========================================")
    print("TUTOR REPORT")
    print("TUTOR:", tutor_name)
    print("DEPARTMENT:", department)
    print("CLASS YEAR:", class_year)
    print("SECTION:", section)
    print("TOTAL STUDENTS:", total_students)
    print("TOTAL CERTIFICATES:", total_certificates)
    print("CERTIFICATE COUNT:", len(certificates))
    print("========================================")


    cursor.close()


    # =========================================================
    # Send to Template
    # =========================================================

    return render_template(
        "tutor/report.html",
        tutor=tutor,
        total_students=total_students,
        total_certificates=total_certificates,
        department_report=department_report,
        year_report=year_report,
        certificates=certificates
    )
#--------------------
@app.route("/view_certificate/<filename>")
def view_certificate(filename):

    if session.get("role") != "tutor":
        return redirect("/")

    cert_folder = os.path.join(
        app.root_path,
        "static",
        "uploads",
        "certificates"
    )

    file_path = os.path.join(cert_folder, filename)

    if not os.path.isfile(file_path):
        return "Certificate file not found", 404

    return send_file(
        file_path,
        as_attachment=False
    )
#-------------------------------------

@app.route("/tutor_export_excel")
def tutor_export_excel():

    if session.get("role") != "tutor":
        return redirect("/")

    cursor = mysql.connection.cursor()

    # ==========================
    # GET TUTOR DETAILS
    # ==========================

    cursor.execute("""
        SELECT
            tutor_name,
            department,
            class_year,
            section
        FROM tutors
        WHERE user_id=%s
    """, (session["user_id"],))

    tutor = cursor.fetchone()

    if not tutor:
        cursor.close()
        flash("Tutor not assigned.")
        return redirect("/tutor")

    tutor_name = tutor[0]
    department = tutor[1]
    class_year = tutor[2]
    section = tutor[3]

    # ==========================
    # GET CERTIFICATE RECORDS
    # ==========================

    if class_year is not None and section is not None:

        cursor.execute("""
            SELECT
                students.register_no,
                students.student_name,
                students.department,
                students.year,
                students.section,
                certificates.certificate_title,
                certificate_categories.category_name,
                certificates.achievement,
                certificates.upload_date

            FROM certificates

            INNER JOIN students
                ON certificates.student_id = students.student_id

            LEFT JOIN certificate_categories
                ON certificates.category_id = certificate_categories.category_id

            WHERE students.year=%s
            AND students.section=%s

            ORDER BY students.student_name
        """, (class_year, section))

    else:

        cursor.execute("""
            SELECT
                students.register_no,
                students.student_name,
                students.department,
                students.year,
                students.section,
                certificates.certificate_title,
                certificate_categories.category_name,
                certificates.achievement,
                certificates.upload_date

            FROM certificates

            INNER JOIN students
                ON certificates.student_id = students.student_id

            LEFT JOIN certificate_categories
                ON certificates.category_id = certificate_categories.category_id

            WHERE students.department=%s

            ORDER BY students.student_name
        """, (department,))

    records = cursor.fetchall()

    print("Tutor:", tutor_name)
    print("Department:", department)
    print("Class Year:", class_year)
    print("Section:", section)
    print("Total Records:", len(records))

    cursor.close()

    # ==========================
    # CREATE EXCEL
    # ==========================

    workbook = Workbook()

    sheet = workbook.active
    sheet.title = "Tutor Report"

    headers = [
        "Register No",
        "Student Name",
        "Department",
        "Year",
        "Section",
        "Certificate",
        "Category",
        "Achievement",
        "Upload Date"
    ]

    sheet.append(headers)

    for row in records:

        upload_date = row[8]

        if upload_date:
            upload_date = upload_date.strftime("%d-%m-%Y")

        sheet.append([
            row[0],
            row[1],
            row[2],
            row[3],
            row[4],
            row[5],
            row[6],
            row[7],
            upload_date
        ])

    output = BytesIO()

    workbook.save(output)

    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name="Tutor_Report.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

#--------------------------
@app.route("/tutor_export_pdf")
def tutor_export_pdf():

    # ==============================
    # TUTOR LOGIN CHECK
    # ==============================

    if session.get("role") != "tutor":
        return redirect("/")


    cursor = mysql.connection.cursor()


    # ==============================
    # GET TUTOR DETAILS
    # ==============================

    cursor.execute("""
        SELECT
            tutor_name,
            department,
            class_year,
            section
        FROM tutors
        WHERE user_id = %s
    """, (session.get("user_id"),))

    tutor = cursor.fetchone()


    if not tutor:

        cursor.close()

        flash("Tutor details not found.")

        return redirect("/tutor")


    tutor_name = tutor[0]
    department = tutor[1]
    class_year = tutor[2]
    section = tutor[3]


    # ==============================
    # GET CERTIFICATE RECORDS
    # ==============================

    # IMPORTANT:
    # If class_year and section are available,
    # filter using year + section.
    #
    # If they are NULL, use department
    # temporarily so certificates are shown.

    if class_year is not None and section is not None:

        cursor.execute("""
            SELECT
                students.register_no,
                students.student_name,
                students.department,
                students.year,
                certificates.certificate_title,
                certificate_categories.category_name,
                certificates.achievement,
                certificates.upload_date

            FROM certificates

            INNER JOIN students
                ON certificates.student_id = students.student_id

            LEFT JOIN certificate_categories
                ON certificates.category_id =
                   certificate_categories.category_id

            WHERE students.year = %s
              AND students.section = %s

            ORDER BY certificates.upload_date DESC
        """, (class_year, section))

    else:

        # ==========================================
        # FALLBACK
        # Tutor class_year / section is NULL
        # So use tutor department
        # ==========================================

        cursor.execute("""
            SELECT
                students.register_no,
                students.student_name,
                students.department,
                students.year,
                certificates.certificate_title,
                certificate_categories.category_name,
                certificates.achievement,
                certificates.upload_date

            FROM certificates

            INNER JOIN students
                ON certificates.student_id = students.student_id

            LEFT JOIN certificate_categories
                ON certificates.category_id =
                   certificate_categories.category_id

            WHERE students.department = %s

            ORDER BY certificates.upload_date DESC
        """, (department,))


    records = cursor.fetchall()


    cursor.close()


    # ==============================
    # CREATE PDF
    # ==============================

    from io import BytesIO

    from reportlab.pdfgen import canvas

    from reportlab.lib.pagesizes import A4

    from reportlab.lib import colors


    buffer = BytesIO()


    pdf = canvas.Canvas(
        buffer,
        pagesize=A4
    )


    width, height = A4


    # ==============================
    # TITLE
    # ==============================

    pdf.setTitle(
        "Tutor Certificate Report"
    )


    pdf.setFont(
        "Helvetica-Bold",
        20
    )

    pdf.drawCentredString(
        width / 2,
        height - 50,
        "STUDENT CERTIFICATE MANAGEMENT SYSTEM"
    )


    pdf.setFont(
        "Helvetica-Bold",
        16
    )

    pdf.drawCentredString(
        width / 2,
        height - 80,
        "Tutor Certificate Report"
    )


    # ==============================
    # TUTOR INFORMATION
    # ==============================

    y = height - 125


    pdf.setFont(
        "Helvetica",
        11
    )


    pdf.drawString(
        50,
        y,
        f"Tutor Name : {tutor_name}"
    )

    y -= 22


    pdf.drawString(
        50,
        y,
        f"Department : {department}"
    )

    y -= 22


    pdf.drawString(
        50,
        y,
        f"Class Year : {class_year if class_year else 'All'}"
    )

    y -= 22


    pdf.drawString(
        50,
        y,
        f"Section : {section if section else 'All'}"
    )

    y -= 35


    # ==============================
    # TOTAL RECORDS
    # ==============================

    pdf.setFont(
        "Helvetica-Bold",
        12
    )


    pdf.drawString(
        50,
        y,
        f"Total Certificate Records : {len(records)}"
    )


    y -= 30


    # ==============================
    # TABLE HEADER
    # ==============================

    pdf.setFont(
        "Helvetica-Bold",
        8
    )


    columns = [
        ("Reg No", 50),
        ("Student", 110),
        ("Department", 190),
        ("Year", 275),
        ("Certificate", 320),
        ("Category", 405),
        ("Achievement", 480),
        ("Date", 555)
    ]


    for title, x in columns:

        pdf.drawString(
            x,
            y,
            title
        )


    y -= 8


    pdf.line(
        50,
        y,
        width - 40,
        y
    )


    y -= 18


    # ==============================
    # CERTIFICATE RECORDS
    # ==============================

    pdf.setFont(
        "Helvetica",
        7
    )


    if records:

        for record in records:

            register_no = record[0]
            student_name = record[1]
            student_department = record[2]
            student_year = record[3]
            certificate_title = record[4]
            category_name = record[5]
            achievement = record[6]
            upload_date = record[7]


            # Format date

            if upload_date:

                try:

                    date_text = upload_date.strftime(
                        "%d-%m-%Y"
                    )

                except:

                    date_text = str(upload_date)

            else:

                date_text = "-"


            # Convert values safely

            register_no = str(
                register_no or "-"
            )

            student_name = str(
                student_name or "-"
            )

            student_department = str(
                student_department or "-"
            )

            student_year = str(
                student_year or "-"
            )

            certificate_title = str(
                certificate_title or "-"
            )

            category_name = str(
                category_name or "-"
            )

            achievement = str(
                achievement or "-"
            )


            # Draw row

            pdf.drawString(
                50,
                y,
                register_no[:10]
            )


            pdf.drawString(
                110,
                y,
                student_name[:14]
            )


            pdf.drawString(
                190,
                y,
                student_department[:12]
            )


            pdf.drawString(
                275,
                y,
                student_year[:8]
            )


            pdf.drawString(
                320,
                y,
                certificate_title[:14]
            )


            pdf.drawString(
                405,
                y,
                category_name[:12]
            )


            pdf.drawString(
                480,
                y,
                achievement[:12]
            )


            pdf.drawString(
                555,
                y,
                date_text
            )


            y -= 20


            # ==========================
            # NEW PAGE
            # ==========================

            if y < 50:

                pdf.showPage()


                y = height - 50


                pdf.setFont(
                    "Helvetica-Bold",
                    12
                )


                pdf.drawString(
                    50,
                    y,
                    "Tutor Certificate Report - Continued"
                )


                y -= 30


                pdf.setFont(
                    "Helvetica",
                    7
                )


    else:

        pdf.setFont(
            "Helvetica-Bold",
            12
        )


        pdf.drawCentredString(
            width / 2,
            y,
            "No certificate records found."
        )


    # ==============================
    # FOOTER
    # ==============================

    pdf.setFont(
        "Helvetica",
        8
    )


    pdf.drawCentredString(
        width / 2,
        30,
        "© 2026 Student Certificate Management System | Designed By PRADEEP From AIDS"
    )


    # ==============================
    # SAVE PDF
    # ==============================

    pdf.save()


    buffer.seek(0)


    from flask import send_file


    return send_file(
    buffer,
    as_attachment=True,
    download_name="Tutor_Certificate_Report.pdf",
    mimetype="application/pdf"
)
#------------------------------

@app.route("/manage_student")
def manage_student():

    if session.get("role") != "tutor":
        return redirect("/")

    cursor = mysql.connection.cursor()

    # ===========================
    # Get Tutor Details
    # ===========================
    cursor.execute("""
        SELECT
            department,
            class_year,
            section
        FROM tutors
        WHERE user_id=%s
    """, (session["user_id"],))

    tutor = cursor.fetchone()

    if not tutor:
        cursor.close()
        flash("Tutor not assigned.")
        return redirect("/tutor")

    department = tutor[0]
    class_year = tutor[1]
    section = tutor[2]

    # ===========================
    # Case 1
    # Department only
    # Show ALL department students
    # ===========================
    if not class_year and not section:

        cursor.execute("""
            SELECT
                s.student_id,
                s.student_name,
                s.register_no,
                s.department,
                s.year,
                s.section,
                s.email,
                s.phone,
                u.username

            FROM students s

            JOIN users u
            ON s.user_id=u.id

            WHERE s.department=%s

            ORDER BY s.student_name
        """, (department,))

    # ===========================
    # Case 2
    # Department + Year
    # Show all students of that year
    # ===========================
    elif class_year and not section:

        cursor.execute("""
            SELECT
                s.student_id,
                s.student_name,
                s.register_no,
                s.department,
                s.year,
                s.section,
                s.email,
                s.phone,
                u.username

            FROM students s

            JOIN users u
            ON s.user_id=u.id

            WHERE
                s.department=%s
                AND s.year=%s

            ORDER BY s.student_name
        """, (
            department,
            class_year
        ))

    # ===========================
    # Case 3
    # Department + Year + Section
    # ===========================
    else:

        cursor.execute("""
            SELECT
                s.student_id,
                s.student_name,
                s.register_no,
                s.department,
                s.year,
                s.section,
                s.email,
                s.phone,
                u.username

            FROM students s

            JOIN users u
            ON s.user_id=u.id

            WHERE
                s.department=%s
                AND s.year=%s
                AND s.section=%s

            ORDER BY s.student_name
        """, (
            department,
            class_year,
            section
        ))

    students = cursor.fetchall()

    cursor.close()

    return render_template(
        "tutor/manage_student.html",
        students=students,
        department=department,
        class_year=class_year,
        section=section
    )

# ==========================
# Add Student
# ==========================

@app.route("/add_student", methods=["POST"])
def add_student():

    if session.get("role") != "tutor":
        return redirect("/")

    student_name = request.form["student_name"]
    register_no = request.form["register_no"]
    department = request.form["department"]
    year = request.form["year"]
    section = request.form["section"]      # NEW
    email = request.form["email"]
    phone = request.form["phone"]
    username = request.form["username"]
    password = request.form["password"]

    cursor = mysql.connection.cursor()

    # Check username
    cursor.execute(
        "SELECT id FROM users WHERE username=%s",
        (username,)
    )

    if cursor.fetchone():
        cursor.close()
        flash("Username already exists")
        return redirect("/manage_student")

    # Insert into users
    cursor.execute("""
        INSERT INTO users
        (
            username,
            password,
            role
        )
        VALUES
        (
            %s,
            %s,
            'student'
        )
    """, (
        username,
        password
    ))

    mysql.connection.commit()

    user_id = cursor.lastrowid

    # Insert into students
    cursor.execute("""
        INSERT INTO students
        (
            user_id,
            student_name,
            register_no,
            department,
            year,
            section,
            email,
            phone
        )
        VALUES
        (
            %s,
            %s,
            %s,
            %s,
            %s,
            %s,
            %s,
            %s
        )
    """, (
        user_id,
        student_name,
        register_no,
        department,
        year,
        section,
        email,
        phone
    ))

    mysql.connection.commit()

    cursor.close()

    flash("Student Created Successfully")

    return redirect("/manage_student")

@app.route("/edit_student/<int:id>")
def edit_student(id):

    if session.get("role") != "tutor":
        return redirect("/")

    cursor = mysql.connection.cursor()

    cursor.execute("""
        SELECT
            students.student_id,
            students.student_name,
            students.register_no,
            students.department,
            students.year,
            students.section,
            students.email,
            students.phone,
            users.username

        FROM students

        INNER JOIN users
            ON students.user_id = users.id

        WHERE students.student_id=%s
    """, (id,))

    student = cursor.fetchone()

    cursor.close()

    return render_template(
        "tutor/edit_student.html",
        student=student
    )

@app.route("/update_student/<int:id>", methods=["POST"])
def update_student(id):

    if session.get("role") != "tutor":
        return redirect("/")

    student_name = request.form["student_name"]
    register_no = request.form["register_no"]
    department = request.form["department"]
    year = request.form["year"]
    section = request.form["section"]      # NEW
    email = request.form["email"]
    phone = request.form["phone"]
    username = request.form["username"]

    cursor = mysql.connection.cursor()

    # Update Student Details
    cursor.execute("""
        UPDATE students
        SET
            student_name=%s,
            register_no=%s,
            department=%s,
            year=%s,
            section=%s,
            email=%s,
            phone=%s
        WHERE student_id=%s
    """, (
        student_name,
        register_no,
        department,
        year,
        section,
        email,
        phone,
        id
    ))

    # Get user_id
    cursor.execute("""
        SELECT user_id
        FROM students
        WHERE student_id=%s
    """, (id,))

    row = cursor.fetchone()

    if row:

        user_id = row[0]

        # Check duplicate username
        cursor.execute("""
            SELECT id
            FROM users
            WHERE username=%s
            AND id!=%s
        """, (
            username,
            user_id
        ))

        if cursor.fetchone():

            cursor.close()
            flash("Username already exists")
            return redirect(f"/edit_student/{id}")

        # Update username
        cursor.execute("""
            UPDATE users
            SET username=%s
            WHERE id=%s
        """, (
            username,
            user_id
        ))

    mysql.connection.commit()

    cursor.close()

    flash("Student Updated Successfully")

    return redirect("/manage_student")

@app.route("/delete_student/<int:id>")
def delete_student(id):

    if session.get("role") != "tutor":
        return redirect("/")

    cursor = mysql.connection.cursor()

    cursor.execute("""
        SELECT user_id
        FROM students
        WHERE student_id=%s
    """, (id,))

    row = cursor.fetchone()

    if row:

        user_id = row[0]

        cursor.execute("DELETE FROM students WHERE student_id=%s", (id,))
        cursor.execute("DELETE FROM users WHERE id=%s", (user_id,))

        mysql.connection.commit()

    cursor.close()

    flash("Student Deleted Successfully")

    return redirect("/manage_student")

#--
@app.route("/download_all_tutor")
def download_all_tutor():

    if session.get("role") != "tutor":
        return redirect("/")

    search = request.args.get("student", "").strip()

    cursor = mysql.connection.cursor()

    # Tutor Class Details
    cursor.execute("""
        SELECT
            class_year,
            section
        FROM tutors
        WHERE user_id=%s
    """, (session["user_id"],))

    tutor = cursor.fetchone()

    if not tutor:
        cursor.close()
        flash("Tutor not found")
        return redirect("/tutor")

    class_year = tutor[0]
    section = tutor[1]

    if search:

        cursor.execute("""
            SELECT certificates.certificate_file

            FROM certificates

            INNER JOIN students
            ON certificates.student_id = students.student_id

            WHERE students.year=%s
            AND students.section=%s
            AND students.student_name LIKE %s
        """, (
            class_year,
            section,
            "%" + search + "%"
        ))

    else:

        cursor.execute("""
            SELECT certificates.certificate_file

            FROM certificates

            INNER JOIN students
            ON certificates.student_id = students.student_id

            WHERE students.year=%s
            AND students.section=%s
        """, (
            class_year,
            section
        ))

    files = cursor.fetchall()

    cursor.close()

    temp = tempfile.NamedTemporaryFile(delete=False, suffix=".zip")

    with zipfile.ZipFile(temp.name, "w") as zipf:

        for row in files:

            filename = row[0]

            filepath = os.path.join(
                app.config["UPLOAD_FOLDER"],
                filename
            )

            if os.path.isfile(filepath):
                zipf.write(filepath, arcname=filename)

    return send_file(
        temp.name,
        as_attachment=True,
        download_name="Tutor_Certificates.zip"
    )
#sreach

@app.route("/search_certificate_tutor")
def search_certificate_tutor():

    if session.get("role") != "tutor":
        return ""

    search = request.args.get("search", "").strip()

    cursor = mysql.connection.cursor()

    # Tutor Details
    cursor.execute("""
        SELECT
            class_year,
            section
        FROM tutors
        WHERE user_id=%s
    """, (session["user_id"],))

    tutor = cursor.fetchone()

    if not tutor:
        cursor.close()
        return ""

    class_year = tutor[0]
    section = tutor[1]

    cursor.execute("""
        SELECT
            students.student_name,
            students.department,
            students.year,
            certificates.certificate_title,
            certificate_categories.category_name,
            certificates.achievement,
            certificates.certificate_file,
            certificates.upload_date,
            certificates.certificate_id

        FROM certificates

        INNER JOIN students
            ON certificates.student_id = students.student_id

        INNER JOIN certificate_categories
            ON certificates.category_id = certificate_categories.category_id

        WHERE
            students.year=%s
            AND students.section=%s
            AND (
                students.student_name LIKE %s
                OR certificates.certificate_title LIKE %s
            )

        ORDER BY certificates.upload_date DESC
    """, (
        class_year,
        section,
        "%" + search + "%",
        "%" + search + "%"
    ))

    data = cursor.fetchall()

    cursor.close()

    html = ""

    for row in data:

        html += f"""
        <tr>

            <td>{row[0]}</td>

            <td>{row[1]}</td>

            <td>{row[2]}</td>

            <td>{row[3]}</td>

            <td>{row[4]}</td>

            <td>
        """

        if row[5] == "Winner":
            html += '<span class="badge bg-success">🏆 Winner</span>'

        elif row[5] == "Runner":
            html += '<span class="badge bg-warning text-dark">🥈 Runner</span>'

        elif row[5] == "Participated":
            html += '<span class="badge bg-primary">🎖 Participated</span>'

        else:
            html += '<span class="badge bg-secondary">📜 Others</span>'

        html += f"""
            </td>

            <td>{row[7]}</td>

            <td>
                <a href="/view/{row[6]}" class="btn btn-primary btn-sm">
                    View
                </a>
            </td>

            <td>
                <a href="/download/{row[6]}" class="btn btn-success btn-sm">
                    Download
                </a>
            </td>

            <td>
                <a href="/delete_certificate/{row[8]}"
                   class="btn btn-danger btn-sm"
                   onclick="return confirm('Delete this certificate?')">
                    Delete
                </a>
            </td>

        </tr>
        """

    return html
#---------------------------
#upload profile
#---------------------------

@app.route("/upload_tutor_photo", methods=["POST"])
def upload_tutor_photo():

    if session.get("role") != "tutor":
        return redirect("/")

    file = request.files["profile_photo"]

    if not file or file.filename == "":
        flash("Please select a photo")
        return redirect("/tutor_profile")

    filename = secure_filename(file.filename)
    file.save(os.path.join(app.config["PROFILE_FOLDER"], filename))

    cursor = mysql.connection.cursor()

    # tutors table
    cursor.execute("""
        UPDATE tutors
        SET profile_photo=%s
        WHERE user_id=%s
    """, (filename, session["user_id"]))

    # users table
    cursor.execute("""
        UPDATE users
        SET profile_photo=%s
        WHERE id=%s
    """, (filename, session["user_id"]))

    mysql.connection.commit()
    cursor.close()

    flash("Profile Photo Updated Successfully")

    return redirect("/tutor_profile")

# ==========================
# Certificate Categories
# ==========================

@app.route("/certificate_categories")
def certificate_categories():

    if session.get("role") != "tutor":
        return redirect("/")

    cursor = mysql.connection.cursor()

    # Tutor Details
    cursor.execute("""
        SELECT
            class_year,
            section
        FROM tutors
        WHERE user_id=%s
    """, (session["user_id"],))

    tutor = cursor.fetchone()

    if not tutor:
        cursor.close()
        flash("Tutor not assigned.")
        return redirect("/tutor")

    class_year = tutor[0]
    section = tutor[1]

    # Category-wise Certificate Count
    cursor.execute("""
        SELECT
            cc.category_id,
            cc.category_name,
            COUNT(c.certificate_id) AS total_count

        FROM certificate_categories cc

        LEFT JOIN certificates c
            ON cc.category_id = c.category_id

        LEFT JOIN students s
            ON c.student_id = s.student_id

        WHERE
            (
                s.year=%s
                AND s.section=%s
            )
            OR s.student_id IS NULL

        GROUP BY
            cc.category_id,
            cc.category_name

        ORDER BY
            cc.category_name
    """, (
        class_year,
        section
    ))

    categories = cursor.fetchall()

    cursor.close()

    return render_template(
        "tutor/certificate_categories.html",
        categories=categories
    )
# ==========================
# Add Category
# ==========================

@app.route("/add_category", methods=["POST"])
def add_category():

    if session.get("role") != "tutor":
        return redirect("/")

    category_name = request.form["category_name"]

    cursor = mysql.connection.cursor()

    cursor.execute("""
        INSERT INTO certificate_categories
        (category_name)
        VALUES (%s)
    """, (category_name,))

    mysql.connection.commit()

    cursor.close()

    flash("Category Added Successfully")

    return redirect("/certificate_categories")


# ==========================
# Edit Category
# ==========================
@app.route("/edit_category/<int:id>")
def edit_category():

    if session.get("role") != "tutor":
        return redirect("/")

    cursor = mysql.connection.cursor()

    cursor.execute("""
        SELECT
            category_id,
            category_name
        FROM certificate_categories
        WHERE category_id=%s
    """, (id,))

    category = cursor.fetchone()

    cursor.close()

    return render_template(
        "tutor/edit_category.html",
        category=category
    )

# ==========================
# Update Category
# ==========================


@app.route("/update_category/<int:id>", methods=["POST"])
def update_category(id):

    if session.get("role") != "tutor":
        return redirect("/")

    category_name = request.form["category_name"]

    cursor = mysql.connection.cursor()

    cursor.execute("""
        UPDATE certificate_categories
        SET category_name=%s
        WHERE category_id=%s
    """, (
        category_name,
        id
    ))

    mysql.connection.commit()

    cursor.close()

    flash("Category Updated Successfully")

    return redirect("/certificate_categories")

# ==========================
# Delete Category
# ==========================

@app.route("/delete_category/<int:id>")
def delete_category(id):

    if session.get("role") != "tutor":
        return redirect("/")

    cursor = mysql.connection.cursor()

    cursor.execute("""
        DELETE FROM certificate_categories
        WHERE category_id=%s
    """, (id,))

    mysql.connection.commit()

    cursor.close()

    flash("Category Deleted Successfully")

    return redirect("/certificate_categories")
#---------------------------

@app.route("/upload_students", methods=["POST"])
def upload_students():

    if session.get("role") != "tutor":
        return redirect("/")

    file = request.files["excel_file"]

    if file.filename == "":
        flash("Please select an Excel file")
        return redirect("/manage_student")

    if not file.filename.lower().endswith(".xlsx"):
        flash("Please upload only .xlsx file")
        return redirect("/manage_student")

    filename = secure_filename(file.filename)
    filepath = os.path.join(UPLOAD_FOLDER, filename)
    file.save(filepath)

    # Read Excel Sheet
    df = pd.read_excel(
        filepath,
        sheet_name="studentss",
        engine="openpyxl"
    )

    print(df.columns.tolist())

    cursor = mysql.connection.cursor()

    for _, row in df.iterrows():

        username = row["Username"]
        password = row["Password"]

        # Check Username
        cursor.execute(
            "SELECT id FROM users WHERE username=%s",
            (username,)
        )

        existing = cursor.fetchone()

        if existing:

            user_id = existing[0]

        else:

            cursor.execute("""
                INSERT INTO users
                (
                    username,
                    password,
                    role
                )
                VALUES
                (
                    %s,
                    %s,
                    'student'
                )
            """, (
                username,
                password
            ))

            mysql.connection.commit()

            user_id = cursor.lastrowid

        # Check Register Number
        cursor.execute(
            """
            SELECT student_id
            FROM students
            WHERE register_no=%s
            """,
            (row["Register Number"],)
        )

        if cursor.fetchone():
            continue

        # Insert Student
        cursor.execute("""
            INSERT INTO students
            (
                user_id,
                register_no,
                student_name,
                department,
                year,
                section,
                email,
                phone
            )
            VALUES
            (
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s
            )
        """, (
            user_id,
            row["Register Number"],
            row["Student Name"],
            row["Department"],
            row["Year"],
            row["Section"],
            row["Email Address"],
            row["Phone Number"]
        ))

    mysql.connection.commit()

    cursor.close()

    flash("Students Imported Successfully")

    return redirect("/manage_student")

#----------------------------

@app.route("/download_sample_excel")
def download_sample_excel():

    sample_path = os.path.join(app.root_path, "static", "sample")

    return send_from_directory(
        sample_path,
        "sample_certificate.xlsx",
        as_attachment=True
    )

#-----------------------------


@app.route("/export_zip")
def export_zip():

    if session.get("role") != "tutor":
        return redirect("/")

    # Search box value
    student_name = request.args.get("student_name", "").strip()

    cursor = mysql.connection.cursor()

    # =========================================================
    # Get Tutor Details
    # =========================================================

    cursor.execute("""
        SELECT
            department,
            class_year,
            section
        FROM tutors
        WHERE user_id=%s
    """, (session["user_id"],))

    tutor = cursor.fetchone()

    if not tutor:
        cursor.close()
        flash("Tutor not assigned.")
        return redirect("/tutor")

    department = tutor[0]
    class_year = tutor[1]
    section = tutor[2]

    # =========================================================
    # Get Certificate Files
    # =========================================================

    if class_year is None or section is None:

        # -----------------------------------------------------
        # No year / section
        # Department-wise certificates
        # -----------------------------------------------------

        if student_name:

            cursor.execute("""
                SELECT
                    certificates.certificate_file
                FROM certificates

                INNER JOIN students
                    ON certificates.student_id =
                       students.student_id

                WHERE students.department=%s
                AND students.student_name LIKE %s
            """, (
                department,
                "%" + student_name + "%"
            ))

        else:

            # Empty search = ALL department certificates

            cursor.execute("""
                SELECT
                    certificates.certificate_file
                FROM certificates

                INNER JOIN students
                    ON certificates.student_id =
                       students.student_id

                WHERE students.department=%s
            """, (department,))

    else:

        # -----------------------------------------------------
        # Year + Section assigned
        # -----------------------------------------------------

        if student_name:

            cursor.execute("""
                SELECT
                    certificates.certificate_file
                FROM certificates

                INNER JOIN students
                    ON certificates.student_id =
                       students.student_id

                WHERE students.department=%s
                AND students.year=%s
                AND students.section=%s
                AND students.student_name LIKE %s
            """, (
                department,
                class_year,
                section,
                "%" + student_name + "%"
            ))

        else:

            # Empty search = ALL certificates
            # from tutor's assigned class

            cursor.execute("""
                SELECT
                    certificates.certificate_file
                FROM certificates

                INNER JOIN students
                    ON certificates.student_id =
                       students.student_id

                WHERE students.department=%s
                AND students.year=%s
                AND students.section=%s
            """, (
                department,
                class_year,
                section
            ))

    certificates = cursor.fetchall()

    cursor.close()

    # =========================================================
    # No Certificates
    # =========================================================

    if not certificates:
        flash("No certificates found.")
        return redirect("/tutor")

    # =========================================================
    # Certificate Folder
    # =========================================================

    cert_folder = os.path.join(
        app.root_path,
        "static",
        "uploads",
        "certificates"
    )

    # =========================================================
    # Create Temporary ZIP
    # =========================================================

    temp_zip = tempfile.NamedTemporaryFile(
        delete=False,
        suffix=".zip"
    )

    zip_path = temp_zip.name
    temp_zip.close()

    # =========================================================
    # Add Only Selected Certificate Files
    # =========================================================

    added_files = 0

    with zipfile.ZipFile(
        zip_path,
        "w",
        zipfile.ZIP_DEFLATED
    ) as zipf:

        for certificate in certificates:

            certificate_file = certificate[0]

            if not certificate_file:
                continue

            file_path = os.path.join(
                cert_folder,
                certificate_file
            )

            if os.path.isfile(file_path):

                zipf.write(
                    file_path,
                    arcname=os.path.basename(file_path)
                )

                added_files += 1

    # =========================================================
    # No Physical Files Found
    # =========================================================

    if added_files == 0:

        try:
            os.remove(zip_path)
        except:
            pass

        flash("Certificate files not found.")
        return redirect("/tutor")

    # =========================================================
    # ZIP File Name
    # =========================================================

    if student_name:

        safe_name = student_name.replace(" ", "_")

        zip_name = safe_name + "_Certificates.zip"

    else:

        zip_name = "Certificates.zip"

    # =========================================================
    # Download ZIP
    # =========================================================

    return send_file(
        zip_path,
        as_attachment=True,
        download_name=zip_name,
        mimetype="application/zip"
    )

# ==========================
# HOD Dashboard
# ==========================

@app.route("/hod")
def hod():

    if session.get("role") != "hod":
        return redirect("/")

    cursor = mysql.connection.cursor()

    # HOD Details
    cursor.execute("""
    SELECT
        hods.hod_name,
        hods.department,
        users.profile_photo

    FROM hods

    INNER JOIN users
        ON hods.user_id = users.id

    WHERE hods.user_id=%s
    """, (session["user_id"],))

    hod = cursor.fetchone()

    if not hod:
        cursor.close()
        flash("HOD not found")
        return redirect("/")

    department = hod[1]

    # ===============================
    # Total Students
    # ===============================

    cursor.execute("""
        SELECT COUNT(*)
        FROM students
        WHERE department=%s
    """, (department,))
    total_students = cursor.fetchone()[0]

    # ===============================
    # Total Tutors
    # ===============================

    cursor.execute("""
        SELECT COUNT(*)
        FROM tutors
        WHERE department=%s
    """, (department,))
    total_tutors = cursor.fetchone()[0]

    # ===============================
    # Total Certificates
    # ===============================

    cursor.execute("""
        SELECT COUNT(*)

        FROM certificates

        INNER JOIN students
        ON certificates.student_id = students.student_id

        WHERE students.department=%s
    """, (department,))

    total_certificates = cursor.fetchone()[0]

    # ===============================
    # Student List
    # ===============================

    cursor.execute("""
        SELECT
            student_id,
            register_no,
            student_name,
            department,
            year,
            email,
            profile_photo

        FROM students

        WHERE department=%s

        ORDER BY student_name
    """, (department,))

    students = cursor.fetchall()

    # ===============================
    # Tutor List
    # ===============================

    cursor.execute("""
        SELECT
            tutor_name,
            department,
            class_year,
            section,
            email,
            profile_photo

        FROM tutors

        WHERE department=%s

        ORDER BY tutor_name
    """, (department,))

    tutors = cursor.fetchall()

    # ===============================
    # Certificate List
    # ===============================

    cursor.execute("""
        SELECT
            students.student_name,
            students.department,
            students.year,
            certificates.certificate_title,
            certificate_categories.category_name,
            certificates.certificate_file,
            certificates.upload_date

        FROM certificates

        INNER JOIN students
            ON certificates.student_id = students.student_id

        INNER JOIN certificate_categories
            ON certificates.category_id = certificate_categories.category_id

        WHERE students.department=%s

        ORDER BY certificates.upload_date DESC
    """, (department,))

    certificates = cursor.fetchall()

    cursor.close()

    return render_template(
        "hod/dashboard.html",
        hod=hod,
        total_students=total_students,
        total_tutors=total_tutors,
        total_certificates=total_certificates,
        students=students,
        tutors=tutors,
        certificates=certificates
    )
#-----------------------------------------
@app.route("/hod_profile")
def hod_profile():

    if session.get("role") != "hod":
        return redirect("/")

    cursor = mysql.connection.cursor()

    cursor.execute("""
        SELECT
            hods.hod_id,
            hods.hod_name,
            hods.department,
            hods.email,
            hods.phone,
            users.username,
            users.profile_photo

        FROM hods

        INNER JOIN users
            ON hods.user_id = users.id

        WHERE users.id=%s
    """, (session["user_id"],))

    hod = cursor.fetchone()

    cursor.close()

    return render_template(
        "hod/profile.html",
        hod=hod
    )


@app.route("/update_hod_profile", methods=["POST"])
def update_hod_profile():

    if session.get("role") != "hod":
        return redirect("/")

    name = request.form["name"]
    email = request.form["email"]
    phone = request.form["phone"]

    profile_photo = None

    upload_folder = os.path.join("static", "uploads", "profile")

    os.makedirs(upload_folder, exist_ok=True)

    if "profile_photo" in request.files:

        file = request.files["profile_photo"]

        if file.filename != "":

            filename = secure_filename(file.filename)

            file.save(os.path.join(upload_folder, filename))

            profile_photo = filename

    cursor = mysql.connection.cursor()

    cursor.execute("""
        UPDATE hods
        SET
            hod_name=%s,
            email=%s,
            phone=%s
        WHERE user_id=%s
    """, (
        name,
        email,
        phone,
        session["user_id"]
    ))

    if profile_photo:

        cursor.execute("""
            UPDATE users
            SET profile_photo=%s
            WHERE id=%s
        """, (
            profile_photo,
            session["user_id"]
        ))

    mysql.connection.commit()

    cursor.close()

    flash("Profile Updated Successfully")

    return redirect("/hod_profile")

#--------------------------
# HOD Reports
# ==========================

@app.route("/hod_reports")
def hod_reports():

    if session.get("role") != "hod":
        return redirect("/")

    cursor = mysql.connection.cursor()

    # ==============================
    # HOD DETAILS
    # ==============================

    cursor.execute("""
        SELECT
            hods.hod_name,
            hods.department,
            hods.email,
            users.profile_photo

        FROM hods

        INNER JOIN users
            ON hods.user_id = users.id

        WHERE users.id = %s
    """, (session["user_id"],))

    hod = cursor.fetchone()

    if not hod:
        cursor.close()
        flash("HOD not assigned.")
        return redirect("/hod")

    department = hod[1]

    # ==============================
    # TOTAL STUDENTS
    # ==============================

    cursor.execute("""
        SELECT COUNT(*)
        FROM students
        WHERE department=%s
    """, (department,))
    total_students = cursor.fetchone()[0]

    # ==============================
    # TOTAL CERTIFICATES
    # ==============================

    cursor.execute("""
        SELECT COUNT(*)
        FROM certificates
        INNER JOIN students
            ON certificates.student_id = students.student_id
        WHERE students.department=%s
    """, (department,))
    total_certificates = cursor.fetchone()[0]

    # ==============================
    # DEPARTMENT REPORT
    # ==============================

    cursor.execute("""
        SELECT department, COUNT(*)
        FROM students
        WHERE department=%s
        GROUP BY department
    """, (department,))
    department_report = cursor.fetchall()

    # ==============================
    # YEAR REPORT
    # ==============================

    cursor.execute("""
        SELECT year, COUNT(*)
        FROM students
        WHERE department=%s
        GROUP BY year
        ORDER BY year
    """, (department,))
    year_report = cursor.fetchall()

    # ==============================
    # CATEGORY REPORT
    # ==============================

    cursor.execute("""
        SELECT
            certificate_categories.category_name,
            COUNT(certificates.certificate_id)

        FROM certificate_categories

        LEFT JOIN certificates
            ON certificate_categories.category_id =
               certificates.category_id

        LEFT JOIN students
            ON certificates.student_id = students.student_id

        WHERE students.department=%s

        GROUP BY certificate_categories.category_name

        ORDER BY certificate_categories.category_name
    """, (department,))

    category_report = cursor.fetchall()

    # ==============================
    # CERTIFICATE REPORT
    # ==============================

    cursor.execute("""
        SELECT
            students.register_no,
            students.student_name,
            students.department,
            students.year,
            certificates.certificate_title,
            certificate_categories.category_name,
            certificates.achievement,
            certificates.certificate_file,
            certificates.upload_date

        FROM certificates

        INNER JOIN students
            ON certificates.student_id = students.student_id

        INNER JOIN certificate_categories
            ON certificates.category_id =
               certificate_categories.category_id

        WHERE students.department=%s

        ORDER BY certificates.upload_date DESC
    """, (department,))

    certificates = cursor.fetchall()

    cursor.close()

    return render_template(
        "hod/report.html",
        hod=hod,
        department=department,
        total_students=total_students,
        total_certificates=total_certificates,
        department_report=department_report,
        year_report=year_report,
        category_report=category_report,
        certificates=certificates
    )

# ==========================
# HOD Export PDF
# ==========================

@app.route("/hod_export_pdf")
def hod_export_pdf():

    if session.get("role") != "hod":
        return redirect("/")

    cursor = mysql.connection.cursor()

    # Get HOD Department
    cursor.execute("""
        SELECT department
        FROM hods
        WHERE user_id=%s
    """, (session["user_id"],))

    hod = cursor.fetchone()

    if not hod:
        cursor.close()
        flash("HOD not assigned.")
        return redirect("/hod")

    department = hod[0]

    # Fetch Certificate Records
    cursor.execute("""
        SELECT
            students.register_no,
            students.student_name,
            students.department,
            students.year,
            certificates.certificate_title,
            certificate_categories.category_name,
            certificates.upload_date

        FROM certificates

        INNER JOIN students
            ON certificates.student_id = students.student_id

        INNER JOIN certificate_categories
            ON certificates.category_id = certificate_categories.category_id

        WHERE students.department=%s

        ORDER BY students.student_name
    """, (department,))

    records = cursor.fetchall()

    cursor.close()

    # Create PDF
    buffer = BytesIO()

    pdf = canvas.Canvas(buffer, pagesize=letter)

    width, height = letter

    y = height - 40

    pdf.setTitle("HOD Certificate Report")

    pdf.setFont("Helvetica-Bold", 16)
    pdf.drawString(180, y, "HOD Certificate Report")

    y -= 30

    pdf.setFont("Helvetica-Bold", 10)

    pdf.drawString(20, y, "Reg No")
    pdf.drawString(85, y, "Student")
    pdf.drawString(180, y, "Dept")
    pdf.drawString(240, y, "Year")
    pdf.drawString(290, y, "Certificate")
    pdf.drawString(430, y, "Category")
    pdf.drawString(510, y, "Date")

    y -= 20

    pdf.setFont("Helvetica", 9)

    for row in records:

        if y < 40:

            pdf.showPage()

            y = height - 40

            pdf.setFont("Helvetica-Bold", 10)

            pdf.drawString(20, y, "Reg No")
            pdf.drawString(85, y, "Student")
            pdf.drawString(180, y, "Dept")
            pdf.drawString(240, y, "Year")
            pdf.drawString(290, y, "Certificate")
            pdf.drawString(430, y, "Category")
            pdf.drawString(510, y, "Date")

            y -= 20

            pdf.setFont("Helvetica", 9)

        pdf.drawString(20, y, str(row[0]))
        pdf.drawString(85, y, str(row[1])[:14])
        pdf.drawString(180, y, str(row[2]))
        pdf.drawString(240, y, str(row[3]))
        pdf.drawString(290, y, str(row[4])[:20])
        pdf.drawString(430, y, str(row[5])[:12])
        pdf.drawString(510, y, str(row[6]))

        y -= 18

    pdf.save()

    buffer.seek(0)

    return send_file(
        buffer,
        as_attachment=True,
        download_name="HOD_Report.pdf",
        mimetype="application/pdf"
    )

# ==========================
# HOD Export Excel
# ==========================

@app.route("/hod_export_excel")
def hod_export_excel():

    if session.get("role") != "hod":
        return redirect("/")

    cursor = mysql.connection.cursor()

    # HOD Department
    cursor.execute("""
        SELECT department
        FROM hods
        WHERE user_id=%s
    """, (session["user_id"],))

    hod = cursor.fetchone()

    if not hod:
        cursor.close()
        flash("HOD Not Found")
        return redirect("/hod")

    department = hod[0]

    # Certificate Records
    cursor.execute("""
        SELECT
            students.register_no,
            students.student_name,
            students.department,
            students.year,
            certificates.certificate_title,
            certificate_categories.category_name,
            certificates.upload_date

        FROM certificates

        INNER JOIN students
            ON certificates.student_id = students.student_id

        INNER JOIN certificate_categories
            ON certificates.category_id = certificate_categories.category_id

        WHERE students.department=%s

        ORDER BY students.student_name
    """, (department,))

    records = cursor.fetchall()

    cursor.close()

    workbook = Workbook()

    sheet = workbook.active

    sheet.title = "HOD Report"

    sheet.append([
        "Register No",
        "Student Name",
        "Department",
        "Year",
        "Certificate",
        "Category",
        "Upload Date"
    ])

    for row in records:
        sheet.append(row)

    output = BytesIO()

    workbook.save(output)

    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name="HOD_Report.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
# ==========================
# Manage Tutors
# ==========================

@app.route("/manage_tutor")
def manage_tutor():

    if session.get("role") != "hod":
        return redirect("/")

    cursor = mysql.connection.cursor()

    # Get HOD Department
    cursor.execute("""
        SELECT department
        FROM hods
        WHERE user_id=%s
    """, (session["user_id"],))

    hod = cursor.fetchone()

    if not hod:
        cursor.close()
        flash("HOD not found")
        return redirect("/logout")

    department = hod[0]

    # Tutor List
    cursor.execute("""
        SELECT
            tutors.tutor_id,
            tutors.tutor_name,
            tutors.department,
            tutors.class_year,
            tutors.section,
            tutors.email,
            users.username

        FROM tutors

        INNER JOIN users
            ON tutors.user_id = users.id

        WHERE tutors.department=%s

        ORDER BY tutors.tutor_name
    """, (department,))

    tutors = cursor.fetchall()

    cursor.close()

    return render_template(
        "hod/manage_tutor.html",
        tutors=tutors,
        department=department
    )

# ==========================
# Add Tutor
# ==========================

@app.route("/add_tutor", methods=["POST"])
def add_tutor():

    if session.get("role") != "hod":
        return redirect("/")

    tutor_name = request.form["tutor_name"]
    department = request.form["department"]
    class_year = request.form["class_year"]
    section = request.form["section"]
    email = request.form["email"]
    username = request.form["username"]
    password = request.form["password"]

    cursor = mysql.connection.cursor()

    # Check Username
    cursor.execute(
        "SELECT id FROM users WHERE username=%s",
        (username,)
    )

    if cursor.fetchone():
        cursor.close()
        flash("Username already exists")
        return redirect("/manage_tutor")

    # Insert into users
    cursor.execute("""
        INSERT INTO users
        (
            username,
            password,
            role
        )
        VALUES
        (
            %s,
            %s,
            'tutor'
        )
    """, (
        username,
        password
    ))

    mysql.connection.commit()

    user_id = cursor.lastrowid

    # Insert into tutors
    cursor.execute("""
        INSERT INTO tutors
        (
            user_id,
            tutor_name,
            department,
            class_year,
            section,
            email
        )
        VALUES
        (
            %s,
            %s,
            %s,
            %s,
            %s,
            %s
        )
    """, (
        user_id,
        tutor_name,
        department,
        class_year,
        section,
        email
    ))

    mysql.connection.commit()

    cursor.close()

    flash("Tutor Created Successfully")

    return redirect("/manage_tutor")
#-------------------------------
#---------------------------------------------------------------------------
@app.route("/edit_tutor/<int:tutor_id>", methods=["GET", "POST"])
def edit_tutor(tutor_id):

    if session.get("role") != "hod":
        return redirect("/")

    cursor = mysql.connection.cursor()

    if request.method == "POST":

        tutor_name = request.form["tutor_name"]
        department = request.form["department"]
        class_year = request.form["class_year"]
        section = request.form["section"]
        email = request.form["email"]
        username = request.form["username"]
        password = request.form["password"]

        # Update tutors table
        cursor.execute("""
            UPDATE tutors
            SET
                tutor_name=%s,
                department=%s,
                class_year=%s,
                section=%s,
                email=%s
            WHERE tutor_id=%s
        """, (
            tutor_name,
            department,
            class_year,
            section,
            email,
            tutor_id
        ))

        # Update username
        cursor.execute("""
            UPDATE users
            SET username=%s
            WHERE id=(
                SELECT user_id
                FROM tutors
                WHERE tutor_id=%s
            )
        """, (
            username,
            tutor_id
        ))

        # Update password (only if entered)
        if password.strip() != "":
            cursor.execute("""
                UPDATE users
                SET password=%s
                WHERE id=(
                    SELECT user_id
                    FROM tutors
                    WHERE tutor_id=%s
                )
            """, (
                password,
                tutor_id
            ))

        mysql.connection.commit()

        cursor.close()

        flash("Tutor Updated Successfully")

        return redirect("/manage_tutor")

    # GET
    cursor.execute("""
        SELECT
            t.tutor_name,
            t.department,
            t.class_year,
            t.section,
            t.email,
            u.username

        FROM tutors t

        INNER JOIN users u
            ON t.user_id = u.id

        WHERE t.tutor_id=%s
    """, (tutor_id,))

    tutor = cursor.fetchone()

    cursor.close()

    return render_template(
        "hod/edit_tutor.html",
        tutor=tutor,
        tutor_id=tutor_id
    )
#-------------------------------------------
@app.route("/delete_tutor/<int:tutor_id>")
def delete_tutor(tutor_id):

    cursor = mysql.connection.cursor()
    cursor.execute("""
    SELECT user_id
    FROM tutors
    WHERE tutor_id=%s
    """, (tutor_id,))

    row = cursor.fetchone()

    if row:

        user_id = row[0]

        # Delete tutor
        cursor.execute("""
        DELETE FROM tutors
        WHERE tutor_id=%s
        """, (tutor_id,))

        # Delete login
        cursor.execute("""
        DELETE FROM users
        WHERE id=%s
        """, (user_id,))

        mysql.connection.commit()

        flash("Tutor deleted successfully")

    cursor.close()

    return redirect("/manage_tutor")
#--------------------
#sreach box
#---------------------

@app.route("/search_certificate")
def search_certificate():
    if session.get("role") != "hod":
        return ""

    search = request.args.get("search", "")

    cursor = mysql.connection.cursor()

    cursor.execute("""
        SELECT
            students.register_no,
            students.student_name,
            certificates.certificate_title,
            certificate_categories.category_name,
            certificates.upload_date,
            certificates.certificate_file

        FROM certificates

        INNER JOIN students
            ON certificates.student_id = students.student_id

        INNER JOIN certificate_categories
            ON certificates.category_id = certificate_categories.category_id

        INNER JOIN hods
            ON hods.department = students.department

        WHERE hods.user_id=%s
AND (
    students.student_name LIKE %s
    OR certificates.certificate_title LIKE %s
)

        ORDER BY students.student_name
    """,  (
    session["user_id"],
    "%" + search + "%",
    "%" + search + "%"
))
    data = cursor.fetchall()

    cursor.close()

    html = ""

    for row in data:
        html += f"""
        <tr>
            <td>{row[0]}</td>
            <td>{row[1]}</td>
            <td>{row[2]}</td>
            <td>{row[3]}</td>
            <td>{row[4]}</td>
            <td>
                <a href="/view/{row[5]}" class="btn btn-success">
                    View
                </a>
            </td>
        </tr>
        """

    return html

#---------------------------
@app.route("/view/<path:filename>")
def view(filename):

    if "role" not in session:
        return redirect("/")

    # =====================================================
    # GOOGLE DRIVE FILE
    # =====================================================

    if (
        filename.startswith("http://")
        or filename.startswith("https://")
    ):

        if (
            "drive.google.com" in filename
            or "drive.usercontent.google.com" in filename
        ):

            import re
            from io import BytesIO
            from googleapiclient.http import MediaIoBaseDownload

            file_id = None

            # /file/d/FILE_ID/view
            match = re.search(
                r"/file/d/([^/]+)",
                filename
            )

            if match:
                file_id = match.group(1)

            # ?id=FILE_ID
            if not file_id:

                match = re.search(
                    r"[?&]id=([^&]+)",
                    filename
                )

                if match:
                    file_id = match.group(1)

            if file_id:

                try:

                    from drive_service import get_drive_service

                    service = get_drive_service()

                    print("================================")
                    print("GOOGLE DRIVE VIEW")
                    print("File ID :", file_id)
                    print("URL     :", filename)
                    print("================================")

                    # Get file information
                    file_info = service.files().get(
                        fileId=file_id,
                        fields="id,name,mimeType"
                    ).execute()

                    original_name = file_info.get(
                        "name",
                        "certificate"
                    )

                    mime_type = file_info.get(
                        "mimeType",
                        "application/octet-stream"
                    )

                    # Get file from Google Drive
                    request = service.files().get_media(
                        fileId=file_id
                    )

                    file_stream = BytesIO()

                    downloader = MediaIoBaseDownload(
                        file_stream,
                        request
                    )

                    done = False

                    while not done:

                        status, done = downloader.next_chunk()

                    file_stream.seek(0)

                    # OPEN IN BROWSER
                    return send_file(
                        file_stream,
                        as_attachment=False,
                        download_name=original_name,
                        mimetype=mime_type
                    )

                except Exception as e:

                    print("================================")
                    print("GOOGLE DRIVE VIEW ERROR")
                    print("File URL :", filename)
                    print("Error    :", str(e))
                    print("================================")

                    return f"Unable to view certificate: {str(e)}", 500

    # =====================================================
    # OLD / LOCAL FILE
    # =====================================================

    return send_from_directory(
        "static/uploads/certificates",
        filename
    )
#-------------------------------


@app.route("/download/<path:filename>")
def download(filename):

    # =====================================================
    # LOGIN CHECK
    # =====================================================

    if "role" not in session:
        return redirect("/")


    # =====================================================
    # GOOGLE DRIVE FILE
    # =====================================================

    if (
        filename.startswith("http://")
        or filename.startswith("https://")
    ):

        # =================================================
        # GOOGLE DRIVE URL CHECK
        # Supports:
        # drive.google.com
        # drive.usercontent.google.com
        # =================================================

        if (
            "drive.google.com" in filename
            or "drive.usercontent.google.com" in filename
        ):

            import re

            file_id = None

            # ---------------------------------------------
            # /file/d/FILE_ID/view
            # ---------------------------------------------

            match = re.search(
                r"/file/d/([^/]+)",
                filename
            )

            if match:
                file_id = match.group(1)

            # ---------------------------------------------
            # ?id=FILE_ID
            # ---------------------------------------------

            if not file_id:

                match = re.search(
                    r"[?&]id=([^&]+)",
                    filename
                )

                if match:
                    file_id = match.group(1)

            # ---------------------------------------------
            # DOWNLOAD FROM GOOGLE DRIVE
            # ---------------------------------------------

            if file_id:

                try:

                    from drive_service import get_drive_service

                    from googleapiclient.http import (
                        MediaIoBaseDownload
                    )

                    from io import BytesIO

                    from flask import send_file

                    # -------------------------------------
                    # Get Google Drive service
                    # -------------------------------------

                    service = get_drive_service()

                    print("================================")
                    print("GOOGLE DRIVE DOWNLOAD")
                    print("File ID :", file_id)
                    print("URL     :", filename)
                    print("================================")

                    # -------------------------------------
                    # Get file information
                    # -------------------------------------

                    file_info = service.files().get(
                        fileId=file_id,
                        fields="id,name,mimeType"
                    ).execute()

                    original_name = file_info.get(
                        "name",
                        "certificate"
                    )

                    mime_type = file_info.get(
                        "mimeType",
                        "application/octet-stream"
                    )

                    print("File Name :", original_name)
                    print("MIME Type :", mime_type)

                    # -------------------------------------
                    # Download file from Google Drive
                    # -------------------------------------

                    request = service.files().get_media(
                        fileId=file_id
                    )

                    file_stream = BytesIO()

                    downloader = MediaIoBaseDownload(
                        file_stream,
                        request
                    )

                    done = False

                    while not done:

                        status, done = downloader.next_chunk()

                        if status:
                            print(
                                "Download Progress :",
                                int(status.progress() * 100),
                                "%"
                            )

                    file_stream.seek(0)

                    # -------------------------------------
                    # FORCE DOWNLOAD
                    # -------------------------------------

                    return send_file(
                        file_stream,
                        as_attachment=True,
                        download_name=original_name,
                        mimetype=mime_type
                    )

                except Exception as e:

                    print("================================")
                    print("GOOGLE DRIVE DOWNLOAD ERROR")
                    print("File URL :", filename)
                    print("Error    :", str(e))
                    print("================================")

                    return f"""
                    <!DOCTYPE html>
                    <html>
                    <head>
                        <title>Download Error</title>
                    </head>

                    <body>

                        <h2>
                            Google Drive Download Failed
                        </h2>

                        <p>
                            <b>File ID:</b> {file_id}
                        </p>

                        <p>
                            <b>Error:</b> {str(e)}
                        </p>

                    </body>
                    </html>
                    """, 500

            else:

                return """
                <h3>Google Drive File ID Not Found</h3>
                <p>
                    The Google Drive URL does not contain
                    a valid file ID.
                </p>
                """, 400


        # =================================================
        # OTHER EXTERNAL URL
        # =================================================

        return redirect(filename)


    # =====================================================
    # OLD LOCAL FILE SUPPORT
    # =====================================================

    folder = os.path.join(
        app.root_path,
        "static",
        "uploads",
        "certificates"
    )

    filename = os.path.basename(filename)

    filepath = os.path.join(
        folder,
        filename
    )

    print("================================")
    print("LOCAL DOWNLOAD ROUTE")
    print("Folder   :", folder)
    print("Filename :", filename)
    print("Full Path:", filepath)
    print("Exists   :", os.path.isfile(filepath))
    print("================================")


    # =====================================================
    # FILE NOT FOUND
    # =====================================================

    if not os.path.isfile(filepath):

        return f"""
        <h3>Certificate File Not Found</h3>

        <p>
            Filename: {filename}
        </p>

        <p>
            Path: {filepath}
        </p>
        """, 404


    # =====================================================
    # LOCAL DOWNLOAD
    # =====================================================

    return send_from_directory(
        folder,
        filename,
        as_attachment=True
    )
#----------------------------
#profile
#-----------------------------

@app.route("/upload_hod_photo", methods=["POST"])
def upload_hod_photo():

    if session.get("role") != "hod":
        return redirect("/")

    file = request.files["profile_photo"]

    if file and file.filename != "":

        filename = secure_filename(file.filename)

        path = os.path.join(
            app.config["UPLOAD_FOLDER_PROFILE"],
            filename
        )

        file.save(path)

        cursor = mysql.connection.cursor()

        cursor.execute("""
            UPDATE hods
            SET profile_photo=%s
            WHERE user_id=%s
        """, (filename, session["user_id"]))

        mysql.connection.commit()

        cursor.close()

        flash("Profile photo uploaded successfully!")

    return redirect("/hod_profile")
#---------------------------
@app.route("/hod_export_zip")
def hod_export_zip():

    if session.get("role") != "hod":
        return redirect("/")

    cursor = mysql.connection.cursor()

    # ==========================================
    # GET LOGGED-IN HOD DEPARTMENT
    # ==========================================

    cursor.execute("""
        SELECT department
        FROM hods
        WHERE user_id=%s
    """, (session["user_id"],))

    hod_data = cursor.fetchone()

    if not hod_data:

        cursor.close()

        flash("HOD not found")

        return redirect("/")


    department = hod_data[0]


    # ==========================================
    # SEARCH TEXT
    # ==========================================

    search_text = request.args.get("search", "").strip()


    # ==========================================
    # EXACT SAME CERTIFICATE SOURCE AS /hod
    # ==========================================

    query = """
        SELECT
            students.student_name,
            students.register_no,
            students.department,
            students.year,
            certificates.certificate_title,
            certificate_categories.category_name,
            certificates.certificate_file,
            certificates.upload_date

        FROM certificates

        INNER JOIN students
            ON certificates.student_id = students.student_id

        INNER JOIN certificate_categories
            ON certificates.category_id =
               certificate_categories.category_id

        WHERE students.department=%s
    """

    params = [department]


    # ==========================================
    # SEARCH FILTER
    # ==========================================

    if search_text:

        query += """
            AND (
                students.student_name LIKE %s
                OR students.register_no LIKE %s
                OR students.department LIKE %s
                OR certificates.certificate_title LIKE %s
                OR certificate_categories.category_name LIKE %s
            )
        """

        search_pattern = "%" + search_text + "%"

        params.extend([
            search_pattern,
            search_pattern,
            search_pattern,
            search_pattern,
            search_pattern
        ])


    # ==========================================
    # ORDER
    # ==========================================

    query += """
        ORDER BY certificates.upload_date DESC
    """


    cursor.execute(query, tuple(params))

    certificates = cursor.fetchall()

    cursor.close()


    # ==========================================
    # NO MATCHING RECORD
    # ==========================================

    if not certificates:

        flash("No certificates found.")

        return redirect("/hod_reports")


    # ==========================================
    # CERTIFICATE FOLDER
    # ==========================================

    cert_folder = os.path.join(
        app.root_path,
        "static",
        "uploads",
        "certificates"
    )


    # ==========================================
    # CREATE ZIP
    # ==========================================

    temp_zip = tempfile.NamedTemporaryFile(
        delete=False,
        suffix=".zip"
    )

    zip_path = temp_zip.name

    temp_zip.close()


    added_files = 0

    used_names = set()


    # ==========================================
    # ADD ONLY MATCHING DATABASE FILES
    # ==========================================

    with zipfile.ZipFile(
        zip_path,
        "w",
        zipfile.ZIP_DEFLATED
    ) as zipf:

        for certificate in certificates:

            # certificate[6] = certificate_file
            filename = certificate[6]

            if not filename:
                continue


            # Prevent unsafe paths
            filename = os.path.basename(filename)


            file_path = os.path.join(
                cert_folder,
                filename
            )


            # ======================================
            # CHECK FILE EXISTS
            # ======================================

            if not os.path.isfile(file_path):
                continue


            # ======================================
            # PREVENT DUPLICATE ZIP NAMES
            # ======================================

            zip_filename = filename

            if zip_filename in used_names:

                name, extension = os.path.splitext(filename)

                counter = 2

                while zip_filename in used_names:

                    zip_filename = (
                        f"{name}_{counter}{extension}"
                    )

                    counter += 1


            used_names.add(zip_filename)


            # ======================================
            # ADD FILE
            # ======================================

            zipf.write(
                file_path,
                arcname=zip_filename
            )

            added_files += 1


    # ==========================================
    # NO FILES FOUND
    # ==========================================

    if added_files == 0:

        if os.path.exists(zip_path):
            os.remove(zip_path)

        flash("Certificate files not found.")

        return redirect("/hod_reports")


    # ==========================================
    # ZIP NAME
    # ==========================================

    if search_text:

        safe_search = "".join(
            c for c in search_text
            if c.isalnum() or c in (" ", "_", "-")
        ).strip()

        if not safe_search:
            safe_search = "Search"

        download_name = (
            safe_search +
            "_Certificates.zip"
        )

    else:

        download_name = (
            "HOD_" +
            department +
            "_Certificates.zip"
        )


    # ==========================================
    # SEND ZIP
    # ==========================================

    return send_file(
        zip_path,
        as_attachment=True,
        download_name=download_name,
        mimetype="application/zip"
    )
# ==========================
# Student Dashboard
# ==========================

@app.route("/student")
def student():

    if session.get("role") != "student":
        return redirect("/")

    cursor = mysql.connection.cursor()

    # =========================================
    # STUDENT DETAILS
    # =========================================

    cursor.execute("""
        SELECT
            student_id,
            user_id,
            register_no,
            student_name,
            department,
            year,
            email,
            phone,
            profile_photo
        FROM students
        WHERE user_id=%s
    """, (session["user_id"],))

    student = cursor.fetchone()

    # =========================================
    # CERTIFICATE CATEGORIES
    # =========================================

    cursor.execute("""
        SELECT
            category_id,
            category_name
        FROM certificate_categories
        ORDER BY category_name
    """)

    categories = cursor.fetchall()

    # =========================================
    # STUDENT CERTIFICATES
    # =========================================

    certificates = []

    if student:

        cursor.execute("""
            SELECT
                certificates.certificate_title,
                certificate_categories.category_name,
                certificates.upload_date,
                certificates.achievement,
                certificates.certificate_file

            FROM certificates

            LEFT JOIN certificate_categories
            ON certificates.category_id =
               certificate_categories.category_id

            WHERE certificates.student_id=%s

            ORDER BY certificates.upload_date DESC
        """, (student[0],))

        certificates = cursor.fetchall()

    cursor.close()

    # =========================================
    # STUDENT DASHBOARD
    # =========================================

    return render_template(
        "student/dashboard.html",
        student=student,
        certificates=certificates,
        categories=categories
    )
#-----------------------------------
@app.route("/student_profile")
def student_profile():

    if session.get("role") != "student":
        return redirect("/")

    cursor = mysql.connection.cursor()

    cursor.execute("""
        SELECT
            student_id,
            register_no,
            student_name,
            department,
            year,
            email,
            phone,
            profile_photo
        FROM students
        WHERE user_id=%s
    """, (session["user_id"],))

    student = cursor.fetchone()

    cursor.close()

    return render_template(
        "student/profile.html",
        student=student
    )
#------------------------------------

@app.route("/update_student_profile", methods=["POST"])
def update_student_profile():

    if session.get("role") != "student":
        return redirect("/")

    student_name = request.form["student_name"]
    email = request.form["email"]
    phone = request.form["phone"]

    cursor = mysql.connection.cursor()

    cursor.execute("""
        UPDATE students
        SET
            student_name=%s,
            email=%s,
            phone=%s
        WHERE user_id=%s
    """, (
        student_name,
        email,
        phone,
        session["user_id"]
    ))

    mysql.connection.commit()

    cursor.close()

    flash("Profile Updated Successfully")

    return redirect("/student_profile")

#--------------------------------


@app.route("/upload_certificate", methods=["POST"])
def upload_certificate():

    if session.get("role") != "student":
        return redirect("/")

    cursor = mysql.connection.cursor()

    try:

        # =========================================
        # GET STUDENT
        # =========================================

        cursor.execute("""
            SELECT student_id
            FROM students
            WHERE user_id=%s
        """, (session["user_id"],))

        student = cursor.fetchone()

        if not student:
            flash("Student not found")
            return redirect("/student")

        student_id = student[0]

        # =========================================
        # GET FORM DATA
        # =========================================

        category_id = request.form.get("category_id")
        certificate_title = request.form.get("certificate_title")
        achievement = request.form.get("achievement")

        file = request.files.get("certificate_file")

        if not file or file.filename == "":
            flash("Please select a certificate")
            return redirect("/student")

        # =========================================
        # FILE NAME
        # =========================================

        filename = secure_filename(file.filename)

        if not filename:
            flash("Invalid certificate filename")
            return redirect("/student")

        # =========================================
        # TEMPORARY LOCAL FOLDER
        # =========================================

        upload_folder = os.path.join(
            "uploads",
            "certificates"
        )

        os.makedirs(
            upload_folder,
            exist_ok=True
        )

        filepath = os.path.join(
            upload_folder,
            filename
        )

        # =========================================
        # SAVE FILE TEMPORARILY
        # =========================================

        file.save(filepath)

        app.logger.info(
            "Certificate saved locally: %s",
            filepath
        )

        # =========================================
        # GOOGLE DRIVE UPLOAD
        # =========================================

        drive_result = upload_to_drive(
            filepath,
            filename
        )

        app.logger.info(
            "Google Drive upload result: %s",
            drive_result
        )

        # =========================================
        # GET DRIVE URL
        # =========================================

        certificate_file = drive_result.get("url")

        if not certificate_file:
            raise Exception(
                "Google Drive URL was not returned."
            )

        app.logger.info(
            "Certificate Drive URL: %s",
            certificate_file
        )

        # =========================================
        # INSERT INTO MYSQL
        # =========================================

        cursor.execute("""
            INSERT INTO certificates
            (
                student_id,
                category_id,
                certificate_title,
                achievement,
                certificate_file,
                upload_date
            )
            VALUES
            (%s, %s, %s, %s, %s, NOW())
        """, (
            student_id,
            category_id,
            certificate_title,
            achievement,
            certificate_file
        ))

        # =========================================
        # CHECK INSERT
        # =========================================

        app.logger.info(
            "Certificate INSERT successful. Row ID: %s",
            cursor.lastrowid
        )

        # =========================================
        # COMMIT
        # =========================================

        mysql.connection.commit()

        app.logger.info(
            "Certificate database commit successful."
        )

        # =========================================
        # DELETE TEMP FILE
        # =========================================

        try:

            if os.path.exists(filepath):
                os.remove(filepath)

        except Exception as e:

            app.logger.warning(
                "Temporary file delete failed: %s",
                e
            )

        flash("Certificate Uploaded Successfully")

        return redirect("/student")

    except Exception as e:

        try:
            mysql.connection.rollback()
        except Exception:
            pass

        app.logger.exception(
            "CERTIFICATE UPLOAD FAILED"
        )

        flash(
            "Certificate upload failed. Please try again."
        )

        return redirect("/student")

    finally:

        cursor.close()
# ==========================
# Download All Certificates
# ==========================


@app.route("/download_all")
def download_all():

    if session.get("role") != "hod":
        return redirect("/")

    search = request.args.get("student", "").strip()

    cursor = mysql.connection.cursor()

    # Get HOD Department
    cursor.execute("""
        SELECT department
        FROM hods
        WHERE user_id=%s
    """, (session["user_id"],))

    hod = cursor.fetchone()

    if not hod:
        cursor.close()
        flash("HOD not found")
        return redirect("/hod")

    department = hod[0]

    # Search student certificates
    if search:

        cursor.execute("""
            SELECT certificates.certificate_file

            FROM certificates

            INNER JOIN students
                ON certificates.student_id = students.student_id

            WHERE students.department=%s
            AND students.student_name LIKE %s
        """, (
            department,
            "%" + search + "%"
        ))

    else:

        cursor.execute("""
            SELECT certificates.certificate_file

            FROM certificates

            INNER JOIN students
                ON certificates.student_id = students.student_id

            WHERE students.department=%s
        """, (department,))

    files = cursor.fetchall()

    cursor.close()

    temp = tempfile.NamedTemporaryFile(delete=False, suffix=".zip")

    with zipfile.ZipFile(temp.name, "w") as zipf:

        for row in files:

            filename = row[0]

            filepath = os.path.join(
                app.config["UPLOAD_FOLDER"],
                filename
            )

            if os.path.isfile(filepath):
                zipf.write(filepath, arcname=filename)

    return send_file(
        temp.name,
        as_attachment=True,
        download_name="Certificates.zip",
        
    )

# ==========================
# Export Excel
# ==========================

@app.route("/export_excel")
def export_excel():

    if session.get("role") != "admin":
        return redirect("/")

    cursor = mysql.connection.cursor()

    cursor.execute("""
        SELECT
            students.register_no,
            students.student_name,
            students.department,
            students.year,
            certificates.certificate_title,
            certificate_categories.category_name,
            certificates.upload_date

        FROM certificates

        INNER JOIN students
            ON certificates.student_id = students.student_id

        INNER JOIN certificate_categories
            ON certificates.category_id = certificate_categories.category_id

        ORDER BY students.student_name
    """)

    records = cursor.fetchall()

    cursor.close()

    workbook = Workbook()

    sheet = workbook.active

    sheet.title = "Certificate Report"

    sheet.append([
        "Register No",
        "Student Name",
        "Department",
        "Year",
        "Certificate",
        "Category",
        "Upload Date"
    ])

    for row in records:
        sheet.append(row)

    output = BytesIO()

    workbook.save(output)

    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name="Certificate_Report.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
# ==========================
# Export PDF
# ==========================

@app.route("/export_pdf")
def export_pdf():

    if session.get("role") != "admin":
        return redirect("/")

    cursor = mysql.connection.cursor()

    cursor.execute("""
        SELECT
            students.register_no,
            students.student_name,
            students.department,
            students.year,
            certificates.certificate_title,
            certificate_categories.category_name,
            certificates.upload_date

        FROM certificates

        INNER JOIN students
            ON certificates.student_id = students.student_id

        INNER JOIN certificate_categories
            ON certificates.category_id = certificate_categories.category_id

        ORDER BY students.student_name
    """)

    records = cursor.fetchall()

    cursor.close()

    buffer = BytesIO()

    pdf = canvas.Canvas(buffer, pagesize=letter)

    width, height = letter

    y = height - 40

    pdf.setTitle("Certificate Report")

    pdf.setFont("Helvetica-Bold", 16)
    pdf.drawString(190, y, "Certificate Report")

    y -= 30

    pdf.setFont("Helvetica-Bold", 10)

    pdf.drawString(20, y, "Reg No")
    pdf.drawString(85, y, "Student")
    pdf.drawString(180, y, "Dept")
    pdf.drawString(240, y, "Year")
    pdf.drawString(290, y, "Certificate")
    pdf.drawString(430, y, "Category")
    pdf.drawString(510, y, "Date")

    y -= 20

    pdf.setFont("Helvetica", 9)

    for row in records:

        if y < 40:
            pdf.showPage()
            y = height - 40
            pdf.setFont("Helvetica", 9)

        pdf.drawString(20, y, str(row[0]))
        pdf.drawString(85, y, str(row[1])[:14])
        pdf.drawString(180, y, str(row[2]))
        pdf.drawString(240, y, str(row[3]))
        pdf.drawString(290, y, str(row[4])[:20])
        pdf.drawString(430, y, str(row[5])[:12])
        pdf.drawString(510, y, str(row[6]))

        y -= 18

    pdf.save()

    buffer.seek(0)

    return send_file(
        buffer,
        as_attachment=True,
        download_name="Certificate_Report.pdf",
        mimetype="application/pdf"
    )

#-------------------------------
#upload profile
#-------------------------------

@app.route("/upload_student_photo", methods=["POST"])
def upload_student_photo():

    if session.get("role") != "student":
        return redirect("/")

    file = request.files["profile_photo"]

    if file.filename == "":
        flash("Please select a photo")
        return redirect("/student")

    filename = secure_filename(file.filename)

    file.save(
        os.path.join(
            app.config["PROFILE_FOLDER"],
            filename
        )
    )

    cursor = mysql.connection.cursor()

    cursor.execute("""
        UPDATE students
        SET profile_photo=%s
        WHERE user_id=%s
    """, (
        filename,
        session["user_id"]
    ))

    mysql.connection.commit()

    cursor.close()

    flash("Profile Photo Updated Successfully")

    return redirect("/student")

# ==========================
# Logout
# ==========================

@app.route("/logout")
def logout():

    session.clear()

    flash("Logged out successfully.")

    return redirect("/")


# ==========================
# Run
# ==========================

if __name__ == "__main__":

    app.run(
        debug=True,
        host="0.0.0.0",
        port=5000
    )

    #---------------
